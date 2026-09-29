import os
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
class ExcelReport:
    def __init__(self,df,dashboard,analysis,recommendations): self.df=df; self.dashboard=dashboard; self.analysis=analysis; self.recommendations=recommendations
    def generate(self,path):
        os.makedirs(os.path.dirname(path),exist_ok=True); wb=Workbook(); self._summary(wb); self._data(wb); self._numeric(wb); self._categorical(wb); self._dates(wb); self._anomalies(wb); self._trends(wb); self._drivers(wb); self._recs(wb)
        for ws in wb.worksheets:
            ws.freeze_panes='A2'
            for col in ws.columns: ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(len(str(c.value or '')) for c in col)+2,40)
        wb.save(path); return path
    def header(self,ws,vals):
        ws.append(vals); fill=PatternFill('solid',fgColor='1D4ED8')
        for c in ws[1]: c.fill=fill; c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(horizontal='center')
    def _summary(self,wb):
        ws=wb.active; ws.title='Executive Summary'; ws['A1']='AI Data Detective Report'; ws['A1'].font=Font(size=18,bold=True); m=self.dashboard.get('primary_metric') or 'Primary Metric'; d=self.dashboard.get('primary_dimension') or 'Primary Dimension'; vals=[('Rows',self.dashboard['rows']),('Columns',self.dashboard['columns']),('Missing Cells',self.dashboard['missing']),('Duplicates',self.dashboard['duplicates']),('Quality Score',self.dashboard['quality']), (f'Total {m}',self.dashboard['primary_metric_total']),(f'Average {m}',self.dashboard['primary_metric_average']),(f'Highest {m}',self.dashboard['primary_metric_max']),(f'Most Common {d}',self.dashboard.get('top_dimension_value')),(f'Highest Total {m} by {d}',self.dashboard.get('top_dimension_metric_value')),('Outlier Count',self.dashboard['outlier_count'])]; ws.append([])
        for a,b in vals: ws.append([a,b]);
        for r in range(3,3+len(vals)): ws[f'A{r}'].font=Font(bold=True)
    def _data(self,wb):
        ws=wb.create_sheet('Cleaned Data'); self.header(ws,list(self.df.columns)); [ws.append(list(r)) for r in self.df.itertuples(index=False)]
    def _numeric(self,wb):
        ws=wb.create_sheet('Numeric Summary'); self.header(ws,['Column','Count','Mean','Median','Min','Max','Sum','Std','Missing','Outliers']); [ws.append([x['column'],x['count'],x['mean'],x['median'],x['min'],x['max'],x['sum'],x['std'],x['missing'],x['outliers']]) for x in self.analysis['numeric_summary']]
    def _categorical(self,wb):
        ws=wb.create_sheet('Categorical Summary'); self.header(ws,['Column','Unique','Missing','Most Common','Frequency']);
        for x in self.analysis['categorical_summary']:
            t=x['top_values'][0] if x['top_values'] else {'value':'','count':0}; ws.append([x['column'],x['unique'],x['missing'],t['value'],t['count']])
    def _dates(self,wb):
        ws=wb.create_sheet('Date Summary'); self.header(ws,['Column','Min Date','Max Date','Unique Dates','Missing']); [ws.append([x['column'],x['min'],x['max'],x['unique_dates'],x['missing']]) for x in self.analysis['datetime_summary']]
    def _anomalies(self,wb):
        ws=wb.create_sheet('Anomalies'); self.header(ws,['Column','Count','Method','Sample Values']); [ws.append([x['column'],x['count'],x['method'],', '.join(map(str,x['sample_values']))]) for x in self.analysis['anomalies']]
    def _trends(self,wb):
        ws=wb.create_sheet('Trends'); self.header(ws,['Column','Direction','Change %','R²','Basis']); [ws.append([x['column'],x['direction'],x['change_percent'],x['r2'],x['basis']]) for x in self.analysis['trends']]
    def _drivers(self,wb):
        ws=wb.create_sheet('Potential Drivers'); self.header(ws,['Metric','Dimension','Segment','Segment Mean','Overall Mean','Difference %','Association Strength','Note']); [ws.append([x['metric'],x['dimension'],x['segment'],x['segment_mean'],x['overall_mean'],x['difference_percent'],x['association_strength'],x['note']]) for x in self.analysis['root_causes']]
    def _recs(self,wb):
        ws=wb.create_sheet('Recommendations'); self.header(ws,['Type','Title','Message']); [ws.append([x['type'],x['title'],x['message']]) for x in self.recommendations]
